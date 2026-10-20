"""Esportazione Excel (scarico) e PDF (regole) del portafoglio LGIM.

Il PDF e' generato da regole.json a ogni esecuzione: testo e parametri non possono divergere dal motore.
"""
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
INTEST = PatternFill("solid", start_color="1F3A5F")


def _foglio(wb, titolo, intestazioni, righe, larghezze, primo=False):
    ws = wb.active if primo else wb.create_sheet()
    ws.title = titolo
    ws.append(intestazioni)
    for r in righe:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10); c.fill = INTEST
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            if isinstance(c.value, float):
                c.number_format = "0.0"
    for i, w in enumerate(larghezze, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def scrivi_xlsx(out, R, path):
    wb = Workbook()
    _foglio(wb, "Portafoglio",
            ["Ticker", "Nome", "ISIN", "Parte", "Area / famiglia", "Peso %", "Target %", "Saldo G", "Saldo S", "Saldo M",
             "Punteggio", "In portafoglio dal", "Dati"],
            [[r["ticker"], r["nome"], r["isin"], r["parte"], r["gruppo"], r["peso"], r["target"], r["D"], r["W"], r["M"],
              r["punteggio"], r["ingresso"], r["dati"]] for r in out["posizioni"]],
            [11, 46, 15, 11, 24, 9, 9, 9, 9, 9, 10, 16, 14], primo=True)
    _foglio(wb, "Candidati",
            ["Ticker", "Nome", "ISIN", "Parte", "Area / famiglia", "Saldo G", "Saldo S", "Saldo M", "Saldo medio", "Punteggio",
             "Ultima barra", "Stato", "Motivo"],
            [[c["ticker"], c["nome"], c["isin"], c["parte"], c["gruppo"], c["D"], c["W"], c["M"], c["saldo_medio"],
              c["punteggio"], c["ultima"], c["stato"], c["motivo"]] for c in out["candidati"]],
            [11, 46, 15, 11, 24, 9, 9, 9, 11, 10, 13, 26, 60])
    mov = []
    for s in out["storico"]:
        if not s["movimenti"]:
            mov.append([s["data"], s["dati_al"], s["azione"], s["motivo"], None, None, None, None, None, None])
        for m in s["movimenti"]:
            mov.append([s["data"], s["dati_al"], s["azione"], s["motivo"], m["ticker"], m["nome"], m["azione"],
                        m["da"], m["a"], m["motivo"]])
    _foglio(wb, "Movimenti",
            ["Data", "Dati al", "Azione", "Motivo dell'azione", "Ticker", "Nome", "Movimento", "Da %", "A %", "Motivo del movimento"],
            mov, [12, 12, 14, 50, 11, 40, 11, 8, 8, 44])
    bm = out["benchmark"]
    _foglio(wb, "Valore del modello", ["Data", "Dati al", "Azione", "Modello"] + bm,
            [[p["data"], p["dati_al"], p["azione"], p["nav"]] + [p.get(b) for b in bm] for p in out["nav_serie"]],
            [12, 12, 16, 11] + [12] * len(bm))
    P = R["parametri"]
    righe = [["Versione regole", R["versione"]], ["Aggiornate il", R["aggiornato"]],
             ["Parte core, peso obiettivo %", P["core_target"]],
             ["Parte satellite, peso obiettivo %", P["satellite"]["target"]],
             ["Banda satellite %", f"{P['satellite']['min']} - {P['satellite']['max']}"],
             ["Tetto per ETF %", P["tetto_peso"]], ["Soglia ribilanciamento settimanale (punti)", P["soglia_ribilanciamento"]],
             ["Stress: quota di peso con saldo G e S negativi %", P["stress_quota_peso"]],
             ["Costo per lato %", P["costo_per_lato_pct"]], ["Paracadute", P["paracadute"]],
             ["Benchmark", ", ".join(P["benchmark"])], [None, None]]
    righe += [[t, x] for t, x in R["testo"]]
    ws = _foglio(wb, "Regole", ["Voce", "Valore"], righe, [46, 120])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# ------------------------------------------------------------------------ PDF
def _sicuro(t):
    """Il PDF usa font standard: niente simboli fuori dal set Latin-1/WinAnsi."""
    for a, b in [("≥", ">="), ("≤", "<="), ("→", "->"), ("–", "-"), ("’", "'"), ("“", '"'), ("”", '"')]:
        t = t.replace(a, b)
    return t


def descr_limiti(regole):
    """Frase in italiano generata dalle regole di limite, cosi' il PDF non diverge dal motore."""
    parti = []
    for r in regole:
        if r.get("valori"):
            v = r["valori"]
            parti.append("massimo %d ETF %s" % (r["max"], ("su " + v[0]) if len(v) == 1 else "sull'insieme " + " e ".join(v)))
    for r in regole:
        if not r.get("valori"):
            nome = {"area": "ogni altra area", "tema": "ogni tema", "famiglia": "ogni famiglia"}.get(r["campo"], r["campo"])
            parti.append("massimo %d per %s" % (r["max"], nome))
    return ", ".join(parti)


def segnaposti(R):
    P = R["parametri"]; pc, ps = P["posizioni"]["core"], P["posizioni"]["satellite"]
    n_core = sum(1 for u in R["universo"] if u["parte"] == "core" and u.get("attivo", True))
    n_sat = sum(1 for u in R["universo"] if u["parte"] == "satellite" and u.get("attivo", True))
    n_ind = len(P["voti"]["indicatori"])
    n_aree = len({u["area"] for u in R["universo"] if u["parte"] == "core" and u.get("attivo", True)})
    return {"n_core": n_core, "n_aree": n_aree, "n_sat": n_sat, "limiti_core": descr_limiti(P["limiti"]["core"]), "limiti_sat": descr_limiti(P["limiti"]["satellite"]),
            "anni": P["download"]["anni_storia"], "salto_max": str(P["download"]["salto_max"]).replace(".", ","),
            "salto_min": str(P["download"]["salto_min"]).replace(".", ","), "rsi_lo": P["soglie"]["rsi"][0],
            "rsi_hi": P["soglie"]["rsi"][1], "adx": P["soglie"]["adx"],
            "n_voti_gs": n_ind + 2 * len(P["voti"]["medie_giornaliero_settimanale"]), "n_voti_m": n_ind + 2 * len(P["voti"]["medie_mensile"]),
            "core_min": pc["min"], "core_max": pc["max"], "core_in": pc["entra_entro"], "core_out": pc["esce_oltre"],
            "sat_pmin": ps["min"], "sat_pmax": ps["max"], "sat_in": ps["entra_entro"], "sat_out": ps["esce_oltre"],
            "perm": P["permanenza_minima_settimane"], "core_t": P["core_target"], "sat_t": P["satellite"]["target"],
            "sat_min": P["satellite"]["min"], "sat_max": P["satellite"]["max"], "tetto": P["tetto_peso"],
            "paracadute": P["paracadute"], "soglia": P["soglia_ribilanciamento"], "stress": P["stress_quota_peso"],
            "costo": str(P["costo_per_lato_pct"]).replace(".", ",")}


def testo_regole(R):
    s = segnaposti(R)
    return [(t, x.format(**s)) for t, x in R["testo"]]


def scrivi_pdf(R, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    titolo = ParagraphStyle("t", parent=st["Title"], fontName="Helvetica-Bold", fontSize=18, alignment=0, spaceAfter=4)
    h = ParagraphStyle("h", parent=st["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=10, spaceAfter=3)
    b = ParagraphStyle("b", parent=st["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13)
    pic = ParagraphStyle("p", parent=b, fontSize=8.5, textColor=colors.HexColor("#555555"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title="Regole del portafoglio LGIM")
    story = [Paragraph("Regole del portafoglio LGIM", titolo),
             Paragraph(escape(_sicuro(f"Versione {R['versione']}, aggiornate il {R['aggiornato']}. Generato automaticamente da regole.json.")), pic)]
    for t, x in testo_regole(R):
        story += [Paragraph(escape(_sicuro(t)), h), Paragraph(escape(_sicuro(x)), b)]
    P = R["parametri"]
    tab = [["Parametro", "Valore"],
           ["Core / satellite (peso obiettivo)", f"{P['core_target']}% / {P['satellite']['target']}%"],
           ["Banda satellite", f"{P['satellite']['min']}% - {P['satellite']['max']}%"],
           ["Tetto per singolo ETF", f"{P['tetto_peso']}%"],
           ["Soglia ribilanciamento settimanale", f"{P['soglia_ribilanciamento']} punti"],
           ["Stress (quota di peso con saldo G e S negativi)", f"{P['stress_quota_peso']}%"],
           ["Permanenza minima", f"{P['permanenza_minima_settimane']} settimane"],
           ["Costo per lato nel valore del modello", f"{str(P['costo_per_lato_pct']).replace('.', ',')}%"],
           ["Paracadute", P["paracadute"]], ["Benchmark", ", ".join(P["benchmark"])]]
    t = Table(tab, colWidths=[9.5 * cm, 7.5 * cm])
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                           ("FONTSIZE", (0, 0), (-1, -1), 9), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story += [Paragraph("Parametri in vigore", h), t, Spacer(1, 8), Paragraph("Universo", h)]
    righe = [["Ticker", "Nome", "Parte", "Area / famiglia"]] + [[u["ticker"], _sicuro(u["nome"]), u["parte"], u.get("area") or u.get("famiglia")]
                                                                  for u in R["universo"]]
    tu = Table(righe, colWidths=[2.2 * cm, 8.6 * cm, 2.2 * cm, 4 * cm], repeatRows=1)
    tu.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 8), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")])]))
    story.append(tu)
    doc.build(story)


def scrivi_xlsx_backtest(bt, R, path):
    wb = Workbook()
    bm = bt["benchmark"]; chiavi = ["modello"] + bm
    st = bt["statistiche"]
    righe = []
    for k in chiavi:
        s = st[k]
        righe.append([k, s["rendimento_totale"], s["rendimento_annuo"], s["volatilita_annua"], s["drawdown_massimo"]])
    ws = _foglio(wb, "Sintesi", ["Serie", "Rendimento totale %", "Rendimento annuo %", "Volatilità annua % (settimanale)", "Drawdown massimo % (settimanale)"],
                 righe, [16, 20, 20, 30, 30], primo=True)
    m = st["modello"]
    extra = [[None] * 5, ["Periodo", f"{bt['periodo']['inizio']} - {bt['periodo']['fine']}"],
             ["Costo per lato %", bt["costo_per_lato_pct"]], ["Decisioni settimanali", m["decisioni"]],
             ["Turnover annuo % del valore", m["turnover_annuo"]], ["Quota media in paracadute %", m["quota_media_paracadute"]],
             ["Quota media nel satellite %", m["quota_media_satellite"]], ["Regole versione", bt["regole_versione"]],
             [None] * 5, ["Limiti", "Universo di oggi (nessun ETF chiuso), prezzi Yahoo Finance, valore rilevato a cadenza settimanale "
                          "(il drawdown è sottostimato). Un buon risultato passato non è una garanzia."]]
    for r in extra:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
    _foglio(wb, "Valore", ["Data", "Dati al", "Azione"] + [("Modello" if k == "modello" else k) for k in chiavi],
            [[p["data"], p["dati_al"], p["azione"]] + [p.get(k) for k in chiavi] for p in bt["serie"]],
            [12, 12, 16] + [12] * len(chiavi))
    _foglio(wb, "Anni solari", ["Anno", "Parziale"] + [("Modello %" if k == "modello" else f"{k} %") for k in chiavi],
            [[a["anno"], "sì" if a["parziale"] else "no"] + [a.get(k) for k in chiavi] for a in bt["anni"]], [10, 10] + [14] * len(chiavi))
    mov = []
    for s in bt["storico"]:
        for mv in s["movimenti"] or [None]:
            mov.append([s["data"], s["dati_al"], s["azione"], s["motivo"]] +
                       ([mv["ticker"], mv["nome"], mv["azione"], mv["da"], mv["a"], mv["motivo"]] if mv else [None] * 6))
    _foglio(wb, "Movimenti", ["Data", "Dati al", "Azione", "Motivo dell'azione", "Ticker", "Nome", "Movimento", "Da %", "A %", "Motivo del movimento"],
            mov, [12, 12, 14, 50, 11, 40, 11, 8, 8, 44])
    nomi = {u["ticker"]: u["nome"] for u in R["universo"]}
    _foglio(wb, "Copertura ETF", ["Ticker", "Nome", "Primo saldo disponibile"],
            [[t, nomi.get(t, t), bt["copertura"].get(t, "mai nel periodo")] for t in nomi], [11, 50, 24])
    wb.save(path)


def scrivi_xlsx_scomposizione(sc, path):
    wb = Workbook()
    chiavi = [k for k in sc["ordine"] + sc["benchmark"] if k in sc["statistiche"]]
    desc = dict(sc["descrizioni"]); desc.update({b: "Benchmark " + b.replace(".MI", "") for b in sc["benchmark"]})
    righe = []
    for k in chiavi:
        s = sc["statistiche"][k]
        righe.append([k, desc.get(k), s["rendimento_totale"], s["rendimento_annuo"], s["volatilita_annua"], s["drawdown_massimo"],
                      s.get("turnover_annuo"), s.get("quota_media_paracadute")])
    ws = _foglio(wb, "Sintesi", ["Serie", "Descrizione", "Rendimento totale %", "Rendimento annuo %", "Volatilità annua %",
                                 "Drawdown massimo %", "Turnover annuo % (senso unico)", "Quota media in paracadute %"],
                 righe, [20, 80, 18, 18, 18, 18, 22, 22], primo=True)
    for r in ([None] * 8, ["Periodo", f"{sc['periodo']['inizio']} - {sc['periodo']['fine']}"],
              ["Costo per lato %", sc["costo_per_lato_pct"]], ["Regole versione", sc["regole_versione"]],
              ["Consigliato L&G valido dal", sc["consigliato_valido_dal"]],
              ["Limiti", "Universo di oggi (nessun ETF chiuso), prezzi Yahoo Finance, valori settimanali: il drawdown è sottostimato."]):
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
    _foglio(wb, "Valore", ["Data"] + chiavi, [[p["data"]] + [p.get(k) for k in chiavi] for p in sc["serie"]], [12] + [16] * len(chiavi))
    _foglio(wb, "Anni solari", ["Anno", "Parziale"] + [k + " %" for k in chiavi],
            [[a["anno"], "sì" if a["parziale"] else "no"] + [a.get(k) for k in chiavi] for a in sc["anni"]], [10, 10] + [18] * len(chiavi))
    _foglio(wb, "Pesi consigliato L&G", ["Ticker", "Peso %"], [[t, w] for t, w in sc["consigliato_pesi"].items()], [14, 10])
    wb.save(path)


# ------------------------------------------------------------ Consigliato LGIM
def _num_it(x):
    return str(x).replace(".", ",")


def testo_consigliato(R):
    C = R["consigliato"]; V = C["variazione"]
    s = {"fonte": C["fonte"], "valido_dal": C["valido_dal"], "min_punti": V["min_punti"], "max_punti": V["max_punti"],
         "coeff": _num_it(V["coefficiente"]), "tetto": R["parametri"]["tetto_peso"]}
    return [(t, x.format(**s)) for t, x in R["testo_consigliato"]]


def scrivi_xlsx_consigliato(cons, R, path):
    wb = Workbook()
    _foglio(wb, "Consigliato",
            ["Ticker", "Nome", "ISIN", "Area / tema", "Peso L&G %", "Peso suggerito %", "Scarto (punti)", "Ultima variazione (punti)",
             "Saldo G", "Saldo S", "Saldo M", "Punteggio", "Dati"],
            [[r["ticker"], r["nome"], r["isin"], r["gruppo"], r["peso_lgim"], r["peso"], r["scarto"], r["ultima_variazione"],
              r["D"], r["W"], r["M"], r["punteggio"], r["dati"]] for r in cons["posizioni"]],
            [11, 46, 15, 20, 11, 14, 12, 16, 9, 9, 9, 10, 14], primo=True)
    mov = []
    for s in cons["storico"]:
        if not s["variazioni"]:
            mov.append([s["data"], s["dati_al"], s["motivo"], s["media_punteggi"], None, None, None, None])
        for t, v in s["variazioni"].items():
            a = s["pesi"].get(t)
            mov.append([s["data"], s["dati_al"], s["motivo"], s["media_punteggi"], t, a - v if a is not None else None, a, v])
    _foglio(wb, "Storico mensile", ["Data", "Dati al", "Motivo", "Media punteggi", "Ticker", "Da %", "A %", "Variazione (punti)"],
            mov, [12, 12, 44, 14, 11, 8, 8, 16])
    V = R["consigliato"]["variazione"]
    righe = [["Fonte dei pesi", cons["fonte"]], ["Pesi L&G validi dal", cons["valido_dal"]], ["Ultimo aggiornamento mensile", cons["aggiornato_il"]],
             ["Prossimo aggiornamento", cons["prossimo_aggiornamento"]], ["Variazione mensile minima (punti)", V["min_punti"]],
             ["Variazione mensile massima (punti)", V["max_punti"]], ["Coefficiente (punti per punto di scarto)", V["coefficiente"]],
             ["Tetto per ETF %", cons["tetto"]], ["Cumulativa", "sì"], ["Versione regole del Consigliato", cons["regole_versione"]], [None, None]]
    righe += [[t, x] for t, x in testo_consigliato(R)]
    ws = _foglio(wb, "Regole", ["Voce", "Valore"], righe, [46, 120])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def scrivi_pdf_consigliato(R, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    titolo = ParagraphStyle("t", parent=st["Title"], fontName="Helvetica-Bold", fontSize=18, alignment=0, spaceAfter=4)
    h = ParagraphStyle("h", parent=st["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=10, spaceAfter=3)
    b = ParagraphStyle("b", parent=st["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13)
    pic = ParagraphStyle("p", parent=b, fontSize=8.5, textColor=colors.HexColor("#555555"))
    C = R["consigliato"]; V = C["variazione"]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title="Regole del Consigliato LGIM")
    story = [Paragraph("Regole del Consigliato LGIM", titolo),
             Paragraph(escape(_sicuro(f"Versione {R.get('versione_consigliato', '1.0')}, pesi L&G validi dal {C['valido_dal']}. "
                                      "Generato automaticamente da regole.json.")), pic)]
    for t, x in testo_consigliato(R):
        story += [Paragraph(escape(_sicuro(t)), h), Paragraph(escape(_sicuro(x)), b)]
    tab = [["Parametro", "Valore"], ["Fonte dei pesi", _sicuro(C["fonte"])], ["Pesi validi dal", C["valido_dal"]],
           ["Frequenza della variazione", V["frequenza"]],
           ["Variazione minima / massima", f"{V['min_punti']} / {V['max_punti']} punti"],
           ["Coefficiente", f"{_num_it(V['coefficiente'])} punti per punto di scarto"],
           ["Tetto per singolo ETF", f"{R['parametri']['tetto_peso']}%"], ["Effetto cumulativo", "sì"]]
    t1 = Table(tab, colWidths=[7 * cm, 10 * cm])
    t1.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 9), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB"))]))
    univ = {u["ticker"]: u for u in R["universo"]}
    righe = [["Ticker", "Nome", "Area / tema", "Peso L&G %"]] + [
        [t, _sicuro(univ[t]["nome"]), univ[t].get("area") or univ[t].get("tema") or univ[t].get("famiglia"), str(w)]
        for t, w in sorted(C["pesi"].items(), key=lambda kv: -kv[1])]
    t2 = Table(righe, colWidths=[2.2 * cm, 8.6 * cm, 3.8 * cm, 2.4 * cm], repeatRows=1)
    t2.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                            ("ALIGN", (3, 0), (3, -1), "RIGHT")]))
    story += [Paragraph("Parametri in vigore", h), t1, Spacer(1, 8), Paragraph("Pesi L&G", h), t2]
    doc.build(story)
